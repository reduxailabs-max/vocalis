"""Tests for vocalis.text.extractors.extract_text.

Uses small real files (built at test time, not committed as binaries)
for every supported format, plus a corrupt file and an image-only PDF
to exercise the error paths for real rather than through a fake.
"""
import logging
from pathlib import Path

import pypdf
import pytest
from docx import Document
from reportlab.pdfgen import canvas

from vocalis.text.extractors import (
    _MAX_TABLE_DEPTH,
    MAX_EXTRACTED_CHARS,
    DocumentOpenError,
    DocumentTooLargeError,
    NoReadableTextError,
    UnsupportedFileTypeError,
    extract_text,
)


def test_extracts_plain_txt(tmp_path: Path) -> None:
    path = tmp_path / "sample.txt"
    path.write_text("Hello from a text file.", encoding="utf-8")
    assert extract_text(path) == "Hello from a text file."


def test_extracts_markdown_without_normalizing_it(tmp_path: Path) -> None:
    path = tmp_path / "sample.md"
    path.write_text("# Heading\n\n**bold**", encoding="utf-8")
    # extractors.py returns raw text; normalization is doc_reader.py's job.
    assert extract_text(path) == "# Heading\n\n**bold**"


def test_extracts_pdf_text(tmp_path: Path) -> None:
    path = tmp_path / "sample.pdf"
    pdf_canvas = canvas.Canvas(str(path), pagesize=(200, 200))
    pdf_canvas.drawString(10, 100, "Hello from a PDF.")
    pdf_canvas.save()
    assert "Hello from a PDF." in extract_text(path)


def test_extracts_docx_text(tmp_path: Path) -> None:
    path = tmp_path / "sample.docx"
    document = Document()
    document.add_paragraph("Hello from a docx file.")
    document.save(str(path))
    assert "Hello from a docx file." in extract_text(path)


def test_extracts_html_and_skips_script_and_style(tmp_path: Path) -> None:
    path = tmp_path / "sample.html"
    path.write_text(
        "<html><head><style>body{color:red}</style></head>"
        "<body><script>alert('x')</script><p>Visible text</p></body></html>",
        encoding="utf-8",
    )
    text = extract_text(path)
    assert "Visible text" in text
    assert "color:red" not in text
    assert "alert" not in text


def test_unsupported_extension_raises() -> None:
    with pytest.raises(UnsupportedFileTypeError):
        extract_text(Path("/tmp/whatever.rtf"))


def test_missing_file_raises_document_open_error(tmp_path: Path) -> None:
    with pytest.raises(DocumentOpenError):
        extract_text(tmp_path / "does_not_exist.txt")


def test_missing_html_file_raises_document_open_error(tmp_path: Path) -> None:
    with pytest.raises(DocumentOpenError):
        extract_text(tmp_path / "does_not_exist.html")


def test_corrupt_pdf_raises_document_open_error(tmp_path: Path) -> None:
    path = tmp_path / "corrupt.pdf"
    path.write_bytes(b"this is not a real pdf file")
    with pytest.raises(DocumentOpenError):
        extract_text(path)


def test_corrupt_docx_raises_document_open_error(tmp_path: Path) -> None:
    path = tmp_path / "corrupt.docx"
    path.write_bytes(b"this is not a real docx file")
    with pytest.raises(DocumentOpenError):
        extract_text(path)


def test_image_only_pdf_raises_no_readable_text_error(tmp_path: Path) -> None:
    path = tmp_path / "scanned.pdf"
    writer = pypdf.PdfWriter()
    writer.add_blank_page(width=200, height=200)
    with path.open("wb") as handle:
        writer.write(handle)
    with pytest.raises(NoReadableTextError):
        extract_text(path)


def test_empty_txt_raises_no_readable_text_error(tmp_path: Path) -> None:
    path = tmp_path / "empty.txt"
    path.write_text("   \n  ", encoding="utf-8")
    with pytest.raises(NoReadableTextError):
        extract_text(path)


def test_oversized_text_raises_document_too_large_error(tmp_path: Path) -> None:
    path = tmp_path / "huge.txt"
    path.write_text("a" * (MAX_EXTRACTED_CHARS + 1), encoding="utf-8")
    with pytest.raises(DocumentTooLargeError):
        extract_text(path)


def test_extracts_docx_table_text_with_no_top_level_paragraphs(tmp_path: Path) -> None:
    path = tmp_path / "table_only.docx"
    document = Document()
    table = document.add_table(rows=1, cols=1)
    table.rows[0].cells[0].text = "Cell text only, no top-level paragraph."
    document.save(str(path))
    assert "Cell text only, no top-level paragraph." in extract_text(path)


def test_extracts_text_from_a_table_nested_inside_a_table_cell(tmp_path: Path) -> None:
    path = tmp_path / "nested.docx"
    document = Document()
    outer = document.add_table(rows=1, cols=1)
    cell = outer.rows[0].cells[0]
    cell.paragraphs[0].text = "Outer cell text."
    inner = cell.add_table(rows=1, cols=1)
    inner.rows[0].cells[0].text = "Inner nested text."
    document.save(str(path))

    text = extract_text(path)

    assert "Outer cell text." in text
    assert "Inner nested text." in text


def test_encrypted_pdf_error_says_it_is_password_protected(tmp_path: Path) -> None:
    path = tmp_path / "locked.pdf"
    writer = pypdf.PdfWriter()
    writer.add_blank_page(width=200, height=200)
    writer.encrypt("secret")
    with path.open("wb") as handle:
        writer.write(handle)

    with pytest.raises(DocumentOpenError, match="password-protected"):
        extract_text(path)


def test_html_file_that_cannot_be_read_raises_document_open_error(tmp_path: Path) -> None:
    path = tmp_path / "folder.html"
    path.mkdir()  # reading a directory raises OSError

    with pytest.raises(DocumentOpenError, match="Could not open folder.html"):
        extract_text(path)


def test_merged_docx_table_cell_text_is_read_only_once(tmp_path: Path) -> None:
    path = tmp_path / "merged.docx"
    document = Document()
    table = document.add_table(rows=1, cols=3)
    cells = table.rows[0].cells
    merged = cells[0].merge(cells[1])
    merged.text = "Merged heading"
    cells[2].text = "Right cell"
    document.save(str(path))

    text = extract_text(path)

    assert text.count("Merged heading") == 1
    assert text.index("Merged heading") < text.index("Right cell")


def test_nested_table_inside_a_merged_cell_is_read_only_once(tmp_path: Path) -> None:
    path = tmp_path / "merged_nested.docx"
    document = Document()
    outer = document.add_table(rows=1, cols=2)
    merged = outer.rows[0].cells[0].merge(outer.rows[0].cells[1])
    merged.paragraphs[0].text = "Merged outer"
    inner = merged.add_table(rows=1, cols=1)
    inner.rows[0].cells[0].text = "Inner once"
    document.save(str(path))

    text = extract_text(path)

    assert text.count("Merged outer") == 1
    assert text.count("Inner once") == 1


def test_cell_merged_across_three_rows_is_read_only_once(tmp_path: Path) -> None:
    path = tmp_path / "vmerged.docx"
    document = Document()
    table = document.add_table(rows=3, cols=2)
    merged = table.cell(0, 0).merge(table.cell(2, 0))
    merged.text = "VMERGED"
    for row_index in range(3):
        table.cell(row_index, 1).text = f"row-{row_index}"
    document.save(str(path))

    text = extract_text(path)

    assert text.count("VMERGED") == 1
    assert text.index("row-0") < text.index("row-1") < text.index("row-2")


def test_side_by_side_merge_is_still_read_once_in_a_multi_row_table(
    tmp_path: Path,
) -> None:
    path = tmp_path / "hmerged.docx"
    document = Document()
    table = document.add_table(rows=2, cols=3)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "HMERGED"
    table.cell(1, 0).text = "below"
    document.save(str(path))

    text = extract_text(path)

    assert text.count("HMERGED") == 1
    assert text.count("below") == 1


def test_different_cells_with_identical_text_are_all_read(tmp_path: Path) -> None:
    path = tmp_path / "same_text.docx"
    document = Document()
    table = document.add_table(rows=2, cols=2)
    for row in table.rows:
        for cell in row.cells:
            cell.text = "same"
    document.save(str(path))

    assert extract_text(path).count("same") == 4


def test_vertical_merge_inside_a_nested_table_is_read_once(tmp_path: Path) -> None:
    path = tmp_path / "nested_vmerge.docx"
    document = Document()
    outer = document.add_table(rows=1, cols=1)
    inner = outer.cell(0, 0).add_table(rows=3, cols=1)
    inner.cell(0, 0).merge(inner.cell(2, 0)).text = "NESTED-V"
    document.save(str(path))

    assert extract_text(path).count("NESTED-V") == 1


def test_warns_when_nested_tables_exceed_the_supported_depth(
    tmp_path: Path, caplog: pytest.LogCaptureFixture
) -> None:
    path = tmp_path / "too_deep.docx"
    document = Document()
    table = document.add_table(rows=1, cols=1)
    # Outer table is depth 0; build well past the limit.
    for level in range(_MAX_TABLE_DEPTH + 3):
        cell = table.rows[0].cells[0]
        cell.paragraphs[0].text = f"level-{level}"
        table = cell.add_table(rows=1, cols=1)
    table.rows[0].cells[0].text = "too-deep-text"
    document.save(str(path))

    with caplog.at_level(logging.WARNING, logger="vocalis"):
        text = extract_text(path)

    assert "level-0" in text
    assert f"level-{_MAX_TABLE_DEPTH}" in text
    assert "too-deep-text" not in text
    warnings = [r.message for r in caplog.records if "nested" in r.message.lower()]
    assert len(warnings) == 1
    assert str(_MAX_TABLE_DEPTH) in warnings[0]
