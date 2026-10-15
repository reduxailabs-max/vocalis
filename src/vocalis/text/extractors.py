"""Extracts raw text from a document file, dispatching by file extension.

Returns the raw text as found in the file — markdown normalization
(normalizer.normalize_markdown) is a separate step performed by
doc_reader.py, so a plain .txt file and a .md file are extracted
identically here.

Every failure mode is a typed exception rather than an unhandled one,
since a corrupt or unusual file the user picks must produce a clear
notification (see doc_reader.py), never crash the popup or the daemon.
Library-specific parsing exceptions (pypdf's, python-docx's) are caught
broadly rather than by exact class, since both libraries can raise a
range of exception types for a malformed file depending on exactly how
it is malformed, and the one thing that matters here is "this file could
not be opened," not which internal exception class reported it.
"""
import html.parser
import logging
from pathlib import Path

import pypdf
from docx import Document as DocxDocument
from docx.table import Table as DocxTable

logger = logging.getLogger("vocalis")

SUPPORTED_EXTENSIONS = frozenset({".txt", ".md", ".pdf", ".docx", ".html", ".htm"})
MAX_EXTRACTED_CHARS = 5_000_000
# Deepest level of tables-inside-tables read from a .docx (sanity limit).
_MAX_TABLE_DEPTH = 10


class UnsupportedFileTypeError(Exception):
    """Raised when the file extension is not one Document Reader supports."""


class DocumentOpenError(Exception):
    """Raised when the file cannot be opened or parsed (corrupt/unreadable)."""


class NoReadableTextError(Exception):
    """Raised when a file opens successfully but contains no extractable
    text (e.g. an image-only scanned PDF with no text layer, or an empty
    file)."""


class DocumentTooLargeError(Exception):
    """Raised when extracted text exceeds MAX_EXTRACTED_CHARS."""


class _HtmlTextExtractor(html.parser.HTMLParser):
    """Collects visible text, skipping <script> and <style> contents."""

    def __init__(self) -> None:
        super().__init__()
        self._skip_depth = 0
        self.text_parts: "list[str]" = []

    def handle_starttag(self, tag: str, attrs: "list[tuple[str, str | None]]") -> None:
        if tag in ("script", "style"):
            self._skip_depth += 1

    def handle_endtag(self, tag: str) -> None:
        if tag in ("script", "style") and self._skip_depth > 0:
            self._skip_depth -= 1

    def handle_data(self, data: str) -> None:
        if self._skip_depth == 0:
            self.text_parts.append(data)


def extract_text(path: Path) -> str:
    """Extract raw text from a document file.

    Args:
        path: Path to the document. Dispatched by suffix (case-insensitive).

    Returns:
        The extracted raw text: non-empty and no larger than
        MAX_EXTRACTED_CHARS.

    Raises:
        UnsupportedFileTypeError: If the extension isn't supported.
        DocumentOpenError: If the file cannot be opened or parsed.
        NoReadableTextError: If extraction succeeds but yields no text.
        DocumentTooLargeError: If extracted text exceeds the size ceiling.
    """
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise UnsupportedFileTypeError(
            f"Unsupported file type {suffix!r}. Supported: "
            f"{', '.join(sorted(SUPPORTED_EXTENSIONS))}"
        )

    if suffix in (".txt", ".md"):
        text = _extract_plain_text(path)
    elif suffix == ".pdf":
        text = _extract_pdf(path)
    elif suffix == ".docx":
        text = _extract_docx(path)
    else:
        text = _extract_html(path)

    if not text.strip():
        raise NoReadableTextError(
            f"No readable text found in {path.name}. If this is a scanned "
            "document, try the OCR Reader feature instead."
        )
    if len(text) > MAX_EXTRACTED_CHARS:
        raise DocumentTooLargeError(
            f"{path.name} contains more than {MAX_EXTRACTED_CHARS:,} characters "
            "of text, which is too large for Document Reader."
        )
    return text


def _read_utf8_text(path: Path) -> str:
    """Read a file as UTF-8 text, replacing undecodable bytes.

    Raises:
        DocumentOpenError: If the file cannot be read.
    """
    try:
        return path.read_bytes().decode("utf-8", errors="replace")
    except OSError as exc:
        logger.warning("Could not open %s: %s", path, exc)
        raise DocumentOpenError(f"Could not open {path.name}: {exc}") from exc


def _extract_plain_text(path: Path) -> str:
    """Read a .txt/.md file as UTF-8 text, replacing undecodable bytes."""
    return _read_utf8_text(path)


def _extract_pdf(path: Path) -> str:
    """Extract text from every page of a PDF file, in page order."""
    try:
        reader = pypdf.PdfReader(str(path))
        if reader.is_encrypted and not reader.decrypt(""):
            raise DocumentOpenError(
                f"{path.name} is password-protected. Remove the password and try again."
            )
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    except DocumentOpenError:
        raise
    except Exception as exc:  # noqa: BLE001 - see module docstring
        logger.warning("Could not open PDF %s: %s", path, exc)
        raise DocumentOpenError(f"Could not open {path.name}: {exc}") from exc


def _extract_docx(path: Path) -> str:
    """Extract text from a .docx file, including table cell text.

    Table text is included because tables are common in real-world
    documents (invoices, resumes, structured reports); a docx whose only
    content is a table would otherwise be misreported as unreadable.
    """
    try:
        document = DocxDocument(str(path))
        parts = [paragraph.text for paragraph in document.paragraphs]
        skipped_deeper_tables = False
        for table in document.tables:
            table_parts, skipped = _docx_table_text(table, depth=0)
            parts.extend(table_parts)
            skipped_deeper_tables = skipped_deeper_tables or skipped
        if skipped_deeper_tables:
            logger.warning(
                "%s has tables nested deeper than the supported limit of %d; "
                "text in the deeper tables was skipped",
                path.name,
                _MAX_TABLE_DEPTH,
            )
        return "\n".join(parts)
    except Exception as exc:  # noqa: BLE001 - see module docstring
        logger.warning("Could not open docx %s: %s", path, exc)
        raise DocumentOpenError(f"Could not open {path.name}: {exc}") from exc


def _docx_table_text(table: "DocxTable", depth: int) -> "tuple[list[str], bool]":
    """Return the paragraph text of every cell in `table`, including text
    in tables nested inside its cells. Nesting deeper than
    _MAX_TABLE_DEPTH is ignored, guarding against pathological documents.

    python-docx's row.cells yields a merged cell once per grid column it
    spans, and a cell merged across rows is yielded again in every row it
    spans. A cell whose underlying XML element was already seen anywhere in
    this table is therefore skipped (order of first appearance is
    preserved). Each nested table keeps its own set of seen cells.

    Returns:
        (paragraph texts, whether any table deeper than _MAX_TABLE_DEPTH
        was skipped)."""
    parts: "list[str]" = []
    skipped_deeper = False
    seen_cells: "set[object]" = set()
    for row in table.rows:
        for cell in row.cells:
            if cell._tc in seen_cells:
                continue
            seen_cells.add(cell._tc)
            parts.extend(paragraph.text for paragraph in cell.paragraphs)
            if depth < _MAX_TABLE_DEPTH:
                for nested in cell.tables:
                    nested_parts, skipped = _docx_table_text(nested, depth + 1)
                    parts.extend(nested_parts)
                    skipped_deeper = skipped_deeper or skipped
            elif cell.tables:
                skipped_deeper = True
    return parts, skipped_deeper


def _extract_html(path: Path) -> str:
    """Extract visible text from an HTML file, skipping script/style content."""
    parser = _HtmlTextExtractor()
    parser.feed(_read_utf8_text(path))
    return "".join(parser.text_parts)
